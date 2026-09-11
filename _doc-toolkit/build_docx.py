"""Render the shared content model to a .docx (also uploads cleanly to Google Docs)."""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, ".")
_src = sys.argv[2] if len(sys.argv) > 2 else 'content'
if _src.endswith('.md'):
    import md_parse
    _mod = md_parse.parse(_src, eyebrow=(sys.argv[3] if len(sys.argv) > 3 else None))
else:
    import importlib
    _mod = importlib.import_module(_src)
TITLE, SUBTITLE, META, BLOCKS = _mod.TITLE, _mod.SUBTITLE, _mod.META, _mod.BLOCKS
EYEBROW = getattr(_mod, 'EYEBROW', 'WAZUH DEPLOYMENT TOOLKIT')

ACCENT   = RGBColor(0x0A, 0x6C, 0x7E)
INK      = RGBColor(0x0F, 0x17, 0x1B)
INK2     = RGBColor(0x51, 0x64, 0x6D)
WARN     = RGBColor(0xA8, 0x5F, 0x09)
CRIT     = RGBColor(0xA3, 0x2A, 0x1C)
SERIF    = "Georgia"
MONO     = "Consolas"

doc = Document()

# ---- page + base styles -------------------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)      # A4
for m in ("top_margin", "bottom_margin"): setattr(sec, m, Inches(0.85))
for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(0.9))

st = doc.styles["Normal"]
st.font.name = SERIF
st.font.size = Pt(10)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.18

def shade(el, hexfill):
    """Apply background shading to a paragraph or table cell."""
    pr = el.get_or_add_pPr() if el.tag.endswith("}p") else el.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexfill)
    pr.append(s)

def border(p, edge="bottom", color="0A6C7E", sz=8):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "3"); e.set(qn("w:color"), color)
    bdr.append(e); pPr.append(bdr)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")
def rich(p, text):
    """Render **bold** and `code` inline markup into runs."""
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = MONO; r.font.size = Pt(8.8)
            r.font.color.rgb = ACCENT
        else:
            p.add_run(part)
    return p

# ---- title block --------------------------------------------------------
p = doc.add_paragraph(); r = p.add_run(EYEBROW)
r.font.name = MONO; r.font.size = Pt(8); r.font.color.rgb = ACCENT; r.bold = True
p.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(); r = p.add_run(TITLE)
r.font.name = MONO; r.font.size = Pt(24); r.bold = True; r.font.color.rgb = INK
p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph(); r = p.add_run(SUBTITLE)
r.font.size = Pt(11); r.font.color.rgb = INK2
p.paragraph_format.space_after = Pt(12)
border(p, "bottom", "0A6C7E", 12)

if META:
    t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [Inches(1.5), Inches(4.4)]
    t.columns[0].width, t.columns[1].width = widths
    for k, v in META:
        row = t.add_row()
        for i, (cell, txt) in enumerate(zip(row.cells, (k, v))):
            cell.width = widths[i]
            cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(1)
            rr = cp.add_run(txt); rr.font.name = MONO; rr.font.size = Pt(8.5)
            rr.font.color.rgb = INK2 if i == 0 else INK
            if i == 0: rr.bold = True
doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ---- block renderers ----------------------------------------------------
def h1(text):
    doc.add_page_break() if False else None
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.name = MONO; r.font.size = Pt(15); r.bold = True; r.font.color.rgb = INK
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(8)
    border(p, "bottom", "0A6C7E", 8)
    p.style = doc.styles["Heading 1"]      # keeps it in the navigation pane / TOC
    for run in p.runs:
        run.font.name = MONO; run.font.size = Pt(15); run.bold = True; run.font.color.rgb = INK

def h2(text):
    p = doc.add_paragraph(style="Heading 2"); r = p.add_run(text)
    r.font.name = MONO; r.font.size = Pt(11.5); r.bold = True; r.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(13); p.paragraph_format.space_after = Pt(4)

def para(text):
    p = doc.add_paragraph(); rich(p, text)

def code(text):
    for i, line in enumerate(text.split("\n")):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0); pf.space_before = Pt(6) if i == 0 else Pt(0)
        pf.left_indent = Inches(0.14); pf.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = MONO; r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor(0x1A, 0x2A, 0x30)
        shade(p._p, "EEF3F4")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def note(kind, label, text):
    colmap = {"info": ("0A6C7E", ACCENT, "E8F3F5"),
              "warn": ("A85F09", WARN,  "FBF2E3"),
              "crit": ("A32A1C", CRIT,  "FAEBE8")}
    hexc, rgbc, fill = colmap[kind]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.12)
    r = p.add_run(label.upper()); r.font.name = MONO; r.font.size = Pt(8)
    r.bold = True; r.font.color.rgb = rgbc
    shade(p._p, fill)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(9); p2.paragraph_format.left_indent = Inches(0.12)
    rich(p2, text)
    for run in p2.runs: run.font.size = Pt(9.5)
    shade(p2._p, fill)
    border(p2, "bottom", hexc, 4)

def table(caption, headers, rows, colw, mono_cols=()):
    total = sum(colw)
    avail = 6.47   # inches inside A4 margins
    widths = [Inches(avail * c / total) for c in colw]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (cell, htxt) in enumerate(zip(t.rows[0].cells, headers)):
        cell.width = widths[i]
        cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.space_before = Pt(2)
        r = cp.add_run(htxt.upper()); r.font.name = MONO; r.font.size = Pt(7.5)
        r.bold = True; r.font.color.rgb = RGBColor(0x33, 0x44, 0x4C)
        shade(cell._tc, "E2EAEC")
    for rowdata in rows:
        cells = t.add_row().cells
        for i, val in enumerate(rowdata):
            cells[i].width = widths[i]
            cp = cells[i].paragraphs[0]
            cp.paragraph_format.space_after = Pt(2); cp.paragraph_format.space_before = Pt(2)
            r = cp.add_run(str(val)); r.font.size = Pt(8.5)
            # Monospace is declared per column in content.py, not guessed per cell.
            if i in mono_cols:
                r.font.name = MONO; r.font.size = Pt(8)
            if str(val) in ("NO", "yes", "no"):
                r.bold = True
                r.font.color.rgb = CRIT if str(val) == "NO" else INK
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def bullets(items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        rich(p, it)
        for run in p.runs: run.font.size = Pt(9.5)

def image(path, caption):
    doc.add_picture(path, width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption); r.font.name = MONO; r.font.size = Pt(7.8); r.font.color.rgb = INK2
    p.paragraph_format.space_after = Pt(10)

DISPATCH = {
    "h1": lambda v: h1(v), "h2": lambda v: h2(v), "h3": lambda v: h2(v),
    "p": lambda v: para(v),
    "code": lambda v: code(v),
    "note": lambda v: note(*v),
    "table": lambda v: table(*v),
    "bullets": lambda v: bullets(v),
    "numbers": lambda v: bullets(v, True),
    "image": lambda v: image(*v),
    "pagebreak": lambda v: doc.add_page_break(),
}
for kind, val in BLOCKS:
    DISPATCH[kind](val)

out = sys.argv[1] if len(sys.argv) > 1 else "Suricata-Wazuh-PoC.docx"
doc.save(out)
print("wrote", out)
